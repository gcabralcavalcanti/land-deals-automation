"""Offline tests: no network, no API key."""
import json
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from land_deals import charts, config, intake, memo, model, synthetic, workbook
from land_deals.llm import LLM, BudgetExceeded, NoRecording
from land_deals.text import parse_number, pmt


@pytest.fixture(scope="module")
def parcels(tmp_path_factory):
    return synthetic.generate(tmp_path_factory.mktemp("p") / "portfolio.json")


@pytest.fixture(scope="module")
def pf(parcels):
    return model.portfolio(parcels)


# ------------------------------------------------------------------------------------- finance math
def test_npv_irr_pmt():
    assert pmt(100000, 0.01, 12) == pytest.approx(8884.88, abs=0.01)
    assert model.npv([-100, 110], 0.0) == pytest.approx(10)
    r = model.monthly_rate(0.12)
    assert model.npv([-100] + [0] * 11 + [112], 0.12) == pytest.approx(0, abs=1e-9)
    assert model.irr([-100] + [0] * 11 + [112]) == pytest.approx(0.12, abs=1e-6)
    assert model.irr([100, 50]) is None
    assert r == pytest.approx(1.12 ** (1 / 12) - 1)


def test_metrics_exposure_and_payback():
    m = model.metrics([0, 1, 2, 3, 4], [-100, -50, 80, 80, 10], 0.0, 0)
    assert m["max_exposure"] == 150 and m["exposure_month"] == 1 and m["payback_month"] == 3
    assert m["profit"] == 20 and m["roi"] == pytest.approx(20 / 150)


# ------------------------------------------------------------------------------------- portfolio
def test_generator(parcels, tmp_path):
    assert len(parcels) == 10 and {p["kind"] for p in parcels} == {"urban", "expansion", "rural"}
    assert synthetic.generate(tmp_path / "x.json") == parcels


def test_strategy_rule(parcels):
    a = config.base()
    for p in parcels:
        s = model.strategy_for(p, a)
        if p["kind"] == "rural" or p["area_m2"] < a["subdivision_min_m2"]:
            assert s == "whole"
    assert any(model.strategy_for(p, a) == "subdivision" for p in parcels)
    p = {**parcels[0], "strategy": "whole"}
    assert model.strategy_for(p, a) == "whole"


def test_flows_add_up(pf):
    for p in pf["projects"]:
        assert sum(p["net"]) == pytest.approx(p["profit"])
        assert sum(p["totals"].values()) == pytest.approx(p["profit"])
        assert all(p["totals"][c] <= 0 for c in model.OUTFLOWS)
        assert all(p["totals"][c] >= 0 for c in model.INFLOWS)
    assert sum(pf["net"]) == pytest.approx(sum(p["profit"] for p in pf["projects"]))
    assert pf["vgv"] == pytest.approx(sum(pf["vgv_by_strategy"].values()))


def test_subdivision_details(parcels):
    p = next(x for x in parcels if model.strategy_for(x, config.base()) == "subdivision")
    r = model.project(p)
    a = config.base()
    assert r["lots"] == int(p["area_m2"] * a["saleable_ratio"] // a["lot_size_m2"])
    assert r["launch"] - r["start"] == a["approval_months"]
    receipts = sum(r["totals"][c] for c in model.INFLOWS)
    assert receipts > r["vgv"]                      # interest and correction on installments
    assert -r["totals"]["sales_tax"] == pytest.approx(receipts * a["sales_tax_rate"])
    no_default = model.project(p, {"default_rate": 0.0})
    assert no_default["totals"]["installments"] > r["totals"]["installments"]


def test_rural_pays_itr_and_car(parcels):
    p = next(x for x in parcels if x["kind"] == "rural")
    r = model.project(p)
    a = config.base()
    expected_taxes = p["acquisition_price"] * (a["itbi_rate"] + a["deed_registry_rate"]) + a["car_service"]
    assert -r["totals"]["acquisition_taxes"] == pytest.approx(expected_taxes)
    monthly_itr = a["itr_rate"] / 12 * p["acquisition_price"]
    assert -r["totals"]["property_tax"] == pytest.approx(monthly_itr * a["hold_months"])


def test_scenarios_and_sensitivity_directions(parcels):
    s = {x["scenario"]: x["npv"] for x in model.scenarios(parcels)}
    assert s["pessimistic"] < s["base"] < s["optimistic"]
    sens = {x["assumption"]: x for x in model.sensitivity(parcels)}
    assert sens["lot_price_factor"]["npv_high"] > sens["lot_price_factor"]["npv_low"]
    assert sens["infra_per_m2"]["npv_high"] < sens["infra_per_m2"]["npv_low"]
    assert sens["discount_rate_year"]["npv_high"] < sens["discount_rate_year"]["npv_low"]
    assert sens["absorption_months"]["npv_high"] < sens["absorption_months"]["npv_low"]


def test_strategy_comparison(parcels):
    rows = model.strategy_comparison(parcels)
    assert rows and all(r["vgv_subdivision"] > r["vgv_whole"] for r in rows)


# ------------------------------------------------------------------------------------- AI layer
class FakeClient:
    def __init__(self, answers):
        self.answers, self.calls = list(answers), 0
        self.messages = self

    def create(self, **kw):
        self.calls += 1
        return SimpleNamespace(stop_reason="end_turn",
                               content=[SimpleNamespace(type="text", text=json.dumps(self.answers.pop(0)))],
                               usage=SimpleNamespace(input_tokens=800, output_tokens=300))


def fake_llm(tmp_path, answers, max_calls=10):
    llm = LLM(mode="demo", max_calls=max_calls, recordings=tmp_path / "rec", log_dir=tmp_path / "logs")
    llm.mode, llm._client = "live", FakeClient(answers)
    return llm


def test_intake_units_and_defaults(tmp_path):
    assert intake.parse_area("12 hectares") == 120000
    assert intake.parse_area("3 alqueires paulistas") == 72600
    assert intake.parse_area("1.800 m²") == 1800
    assert intake.parse_money("R$ 3,8 milhões") == 3_800_000
    assert intake.parse_money("R$ 650 mil") == 650_000
    assert intake.parse_money("R$ 540.000") == 540_000
    text = "Terreno urbano de 1.800 m² no Município F. Compra por R$ 540.000."
    items = [{"field": "name", "value": "Terreno F", "quote": ""}, {"field": "kind", "value": "urban", "quote": ""},
             {"field": "town", "value": "Município F", "quote": "no Município F"},
             {"field": "area", "value": "1.800 m²", "quote": "Terreno urbano de 1.800 m²"},
             {"field": "acquisition_price", "value": "R$ 540.000", "quote": "Compra por R$ 540.000"},
             {"field": "raw_price_m2", "value": "", "quote": ""}, {"field": "lot_price_m2", "value": "", "quote": ""}]
    path = tmp_path / "t.txt"
    path.write_text(text, encoding="utf-8")
    res = intake.run(fake_llm(tmp_path, [{"items": items[:3], "notes": ""}, {"items": items, "notes": ""}]), path)
    assert res["problems"] == []
    assert res["parcel"]["area_m2"] == 1800 and res["parcel"]["raw_price_m2"] == 300
    assert res["source"]["raw_price_m2"] == "default"
    portfolio = tmp_path / "pf.json"
    portfolio.write_text("[]", encoding="utf-8")
    assert intake.accept(res, portfolio)["id"] == "T01"


def test_intake_rejects_invented_quote(tmp_path):
    text = "Sítio de 3 alqueires. Valor: R$ 650 mil."
    items = [{"field": f, "value": "", "quote": ""} for f in intake.FIELDS]
    items[3] = {"field": "area", "value": "5 alqueires", "quote": "Sítio de 5 alqueires"}
    res = intake.to_parcel({"items": items, "notes": ""}, text)
    assert any("area" in p for p in res["problems"])
    with pytest.raises(ValueError):
        intake.accept(res, tmp_path / "pf.json")


def _memo_answer(**kw):
    base = {"recommendation": "invest_with_conditions", "headline": "Subdivision clears the hurdle.",
            "thesis": "The plan works if lots sell on time.", "risks": ["Slower sales", "Infrastructure overrun"],
            "conditions": []}
    return {**base, **kw}


def test_memo_number_check(pf, tmp_path):
    p = max(pf["projects"], key=lambda x: x["npv"])
    fs = memo.facts(p, 0.14)
    irr = fs["IRR (per year)"]
    assert memo.check(_memo_answer(thesis=f"IRR of {irr} beats the hurdle."), fs, p["npv"]) == []
    assert memo.check(_memo_answer(thesis="IRR of 99.9% beats the hurdle."), fs, p["npv"]) == ["number 99.9 is not in FACTS"]
    neg = min(pf["projects"], key=lambda x: x["npv"])
    assert "negative NPV" in " ".join(memo.check(_memo_answer(recommendation="invest"), memo.facts(neg, 0.14), neg["npv"]))


def test_memo_retry_then_reject(pf, tmp_path):
    p = pf["projects"][0]
    bad = _memo_answer(thesis="It returns 77.7% a year.")
    good = _memo_answer(recommendation="do_not_invest")
    ok = memo.write(fake_llm(tmp_path / "a", [bad, good]), p)
    assert ok["status"] == "checked" and ok["recommendation"] == "do_not_invest"
    rejected = memo.write(fake_llm(tmp_path / "b", [bad, bad]), p)
    assert rejected["status"].startswith("rejected")


def test_recordings_and_cap(pf, tmp_path):
    p = pf["projects"][0]
    good = _memo_answer(recommendation="do_not_invest")
    memo.write(fake_llm(tmp_path, [good]), p)
    replay = LLM(mode="demo", recordings=tmp_path / "rec", log_dir=tmp_path / "logs")
    assert memo.write(replay, p)["recommendation"] == "do_not_invest" and replay.calls == 0
    with pytest.raises(NoRecording):
        memo.write(replay, pf["projects"][1])
    with pytest.raises(BudgetExceeded):
        memo.write(fake_llm(tmp_path / "c", [], max_calls=0), p)


# ------------------------------------------------------------------------------------- outputs
def test_outputs(parcels, pf, tmp_path):
    scen, sens, strat = model.scenarios(parcels), model.sensitivity(parcels), model.strategy_comparison(parcels)
    memos = [{"id": p["id"], "status": "not generated", "facts": {}} for p in pf["projects"]]
    xl = workbook.build(tmp_path / "w.xlsx", parcels, pf, scen, sens, strat, memos)
    wb = load_workbook(xl)
    assert wb.sheetnames == ["Summary", "Assumptions", "Parcels", "Costs", "Cash flow", "Sensitivity", "Strategy", "Memos"]
    assert wb["Parcels"].max_row == 11 and wb["Cash flow"].max_row == len(pf["months"]) + 1
    assert wb["Assumptions"].max_row == len(config.ASSUMPTIONS) + 1
    html = charts.dashboard(tmp_path / "d.html", pf, scen, sens, strat, memos).read_text(encoding="utf-8")
    assert "chart.umd.min.js" in html and '"parcels"' in html
    pdf = charts.pdf_report(tmp_path / "r.pdf", pf, scen, sens, strat, memos)
    assert pdf.read_bytes()[:4] == b"%PDF"
    assert parse_number("R$ 1.250.000,00") == 1250000
